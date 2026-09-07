from docx import Document
import pytest

from core.export_manager import ExportManager
from core.logica import ControladorLogico


def test_diary_export_validates_structured_entries(tmp_path):
    with pytest.raises(TypeError, match="lista de diccionarios"):
        ExportManager.export_diary("texto plano", "Proyecto", str(tmp_path / "diario.docx"))
    with pytest.raises(TypeError, match="lista de diccionarios"):
        ExportManager.export_diary(["entrada"], "Proyecto", str(tmp_path / "diario.docx"))


def test_diary_export_writes_content_and_format(tmp_path):
    path = tmp_path / "diario.docx"
    ExportManager.export_diary(
        [{"date": "2026-09-06T14:30:00", "author": "Ana", "message": "Hallazgo central"}],
        "Proyecto",
        str(path),
    )

    document = Document(str(path))
    texts = [paragraph.text for paragraph in document.paragraphs]
    assert texts[0] == "Diario de codificación - Proyecto"
    assert any("Ana - 06/09/2026 a las 14:30" in text for text in texts)
    assert "Hallazgo central" in texts
    header = next(paragraph for paragraph in document.paragraphs if "Ana -" in paragraph.text)
    assert header.runs[0].bold is True


def test_diary_entries_propagate_through_controller_signal(monkeypatch):
    entries = [{"date": "2026-09-06T10:00:00", "author": "L", "message": "Entrada"}]
    received = []
    controller = ControladorLogico()
    controller.export_success.connect(lambda export_type, path: received.append((export_type, path)))
    calls = []
    monkeypatch.setattr(
        ExportManager,
        "export_diary",
        lambda actual_entries, name, path: calls.append((actual_entries, name, path)),
    )

    controller.req_export_diary(entries, "Proyecto", "salida.docx")

    assert calls == [(entries, "Proyecto", "salida.docx")]
    assert received == [("Diario", "salida.docx")]


def test_diary_qt_signal_declares_a_list(main_window):
    signatures = [
        bytes(main_window.metaObject().method(index).methodSignature()).decode()
        for index in range(main_window.metaObject().methodCount())
    ]
    assert "signal_req_export_diary(QVariantList,QString,QString)" in signatures
